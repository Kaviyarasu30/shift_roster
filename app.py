"""
app.py
------
This is the "web server" part of the project. It uses a Python tool
called Flask, which lets a Python script listen for visitors on a
web browser and respond to them.

What this file does:
1. When someone opens the website, it shows them the form (index.html).
2. When they click "Generate Roster", the browser sends the form data
   here. This file reads that data, hands it to roster_logic.py to do
   the actual scheduling work, then builds an Excel file and sends it
   back to the browser as a download.

You do not need to understand every line to use this project — but
the comments explain what each part is for, in case you want to learn.
"""

from datetime import datetime
from io import BytesIO

from flask import Flask, render_template, request, send_file
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from roster_logic import generate_roster, MORNING, EVENING, NIGHT, OFF, LEAVE

app = Flask(__name__)

# Colors used to shade each shift type in the downloaded Excel file,
# so it's easy to scan visually.
SHIFT_COLORS = {
    MORNING: "00B050",   # Green
    EVENING: "FFFF00",   # Yellow
    NIGHT: "FF0000",     # Red
    OFF: "BFBFBF",       # Gray
    LEAVE: "5B9BD5",     # Blue
}


@app.route("/")
def index():
    """Show the empty form when someone visits the site."""
    return render_template("index.html")


@app.route("/generate", methods=["POST"])
def generate():
    """Handle the submitted form: build the roster and return an .xlsx file."""
    form = request.form

    team_name = form.get("team_name", "Team").strip() or "Team"
    start_date = datetime.strptime(form.get("start_date"), "%Y-%m-%d").date()
    end_date = datetime.strptime(form.get("end_date"), "%Y-%m-%d").date()
    if end_date < start_date:
        return "The 'To' date must be on or after the 'From' date.", 400
    if (end_date - start_date).days > 366:
        return "That range is longer than a year — please pick a shorter period.", 400

    # Employee names, chosen shift, AND chosen off-days all arrive as
    # parallel lists — one entry per employee row in the form.
    raw_names = form.getlist("employee_name")
    raw_shifts = form.getlist("employee_shift")
    raw_off_days = form.getlist("employee_off_days")  # comma-separated weekday numbers, e.g. "5,6"
    employees = []
    employee_shifts = {}
    per_employee_off_days = {}
    for i, name in enumerate(raw_names):
        name = name.strip()
        if not name:
            continue
        employees.append(name)
        employee_shifts[name] = raw_shifts[i] if i < len(raw_shifts) else "Morning"
        off_str = raw_off_days[i] if i < len(raw_off_days) else ""
        per_employee_off_days[name] = [int(x) for x in off_str.split(",") if x.strip() != ""]

    if not employees:
        return "Please add at least one employee before generating.", 400

    weekly_off_days = [int(x) for x in form.getlist("weekly_off_days")]

    rules = {
        "weekly_off_enabled": form.get("weekly_off_enabled") == "on",
        "weekly_off_mode": form.get("weekly_off_mode", "rotating"),
        "weekly_off_count": form.get("weekly_off_count", 1),
        "weekly_off_days": weekly_off_days,
        "per_employee_off_days": per_employee_off_days,
        "min_staff_morning": form.get("min_staff_morning", 1),
        "min_staff_evening": form.get("min_staff_evening", 1),
        "min_staff_night": form.get("min_staff_night", 1),
    }

    # ---- Hard check: is the minimum staff rule even POSSIBLE to meet? ----
    # This counts, ignoring leaves/off-days entirely, how many employees
    # are assigned to each shift for the whole period. If that base
    # headcount is already below the minimum you asked for, no amount of
    # rearranging off-days can fix it — so we stop right here with a
    # clear message, instead of quietly generating a roster full of
    # shortage warnings.
    min_staff = {
        MORNING: int(rules["min_staff_morning"]),
        EVENING: int(rules["min_staff_evening"]),
        NIGHT: int(rules["min_staff_night"]),
    }
    assigned_counts = {MORNING: 0, EVENING: 0, NIGHT: 0}
    for emp in employees:
        assigned_counts[employee_shifts[emp]] += 1

    shortfalls = []
    for shift in (MORNING, EVENING, NIGHT):
        if assigned_counts[shift] < min_staff[shift]:
            shortfalls.append({
                "shift": shift,
                "required": min_staff[shift],
                "assigned": assigned_counts[shift],
            })

    if shortfalls:
        return render_template("staffing_error.html", team_name=team_name, shortfalls=shortfalls), 400

    # Leave rows also arrive as parallel lists.
    leave_employees = form.getlist("leave_employee")
    leave_from = form.getlist("leave_from")
    leave_to = form.getlist("leave_to")
    leaves = []
    for emp, f, t in zip(leave_employees, leave_from, leave_to):
        if not emp or not f or not t:
            continue
        leaves.append({
            "employee": emp.strip(),
            "from": datetime.strptime(f, "%Y-%m-%d").date(),
            "to": datetime.strptime(t, "%Y-%m-%d").date(),
        })

    result = generate_roster(team_name, start_date, end_date, employees, employee_shifts, rules, leaves)

    # ---- Hard check #2: did any SPECIFIC day dip below the minimum? ----
    # This catches cases where you have enough people overall, but a
    # leave happens to land on the same day as someone's weekly off,
    # temporarily dropping a shift below the minimum. By default this
    # also blocks the download (instead of just noting it in a
    # spreadsheet tab you might not notice) so short-staffed days never
    # silently ship. Tick "Generate anyway" in the Rules section if you
    # need the file regardless (e.g. an unavoidable emergency leave).
    allow_understaffed = form.get("allow_understaffed_days") == "on"
    if result["warnings"] and not allow_understaffed:
        return render_template(
            "day_shortfall_error.html",
            team_name=team_name,
            issues=result["warnings"],
        ), 400

    excel_bytes = build_excel(team_name, start_date, end_date, result)

    filename = (
        f"{team_name}_Roster_{start_date.isoformat()}_to_{end_date.isoformat()}.xlsx"
    ).replace(" ", "_")
    return send_file(
        excel_bytes,
        as_attachment=True,
        download_name=filename,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


def build_excel(team_name, start_date, end_date, result):
    """Turn the schedule dict into a formatted .xlsx file, returned as bytes."""
    days = result["days"]
    assignments = result["assignments"]
    warnings = result["warnings"]

    wb = Workbook()
    ws = wb.active
    ws.title = "Roster"

    bold = Font(bold=True)
    center = Alignment(horizontal="center", vertical="center")
    thin = Side(style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    date_range_label = f"{start_date.strftime('%d %b %Y')} – {end_date.strftime('%d %b %Y')}"

    # Title row
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(days) + 1)
    title_cell = ws.cell(row=1, column=1, value=f"{team_name} — Shift Roster ({date_range_label})")
    title_cell.font = Font(bold=True, size=14)
    title_cell.alignment = center

    # Header row: Employee | 09 Jun | 10 Jun | ... (full date, since a range
    # can cross month boundaries, plain day numbers like "9" would be
    # ambiguous), with weekday underneath.
    header_row = 3
    ws.cell(row=header_row, column=1, value="Employee").font = bold
    for i, d in enumerate(days, start=2):
        c = ws.cell(row=header_row, column=i, value=d.strftime("%d %b"))
        c.font = Font(bold=True, size=9)
        c.alignment = center
        c.border = border
        weekday_cell = ws.cell(row=header_row + 1, column=i, value=d.strftime("%a"))
        weekday_cell.alignment = center
        weekday_cell.font = Font(size=8, italic=True)

    ws.column_dimensions["A"].width = 20
    for i in range(2, len(days) + 2):
        ws.column_dimensions[get_column_letter(i)].width = 8

    # One row per employee
    data_start_row = header_row + 2
    for row_offset, emp in enumerate(assignments.keys()):
        row = data_start_row + row_offset
        ws.cell(row=row, column=1, value=emp).font = bold
        for col_offset, d in enumerate(days, start=2):
            code = assignments[emp][d]
            label = {MORNING: "A", EVENING: "B", NIGHT: "C", OFF: "WO", LEAVE: "L"}[code]
            cell = ws.cell(row=row, column=col_offset, value=label)
            cell.alignment = center
            cell.border = border
            cell.fill = PatternFill("solid", fgColor=SHIFT_COLORS[code])

    # Legend
    legend_row = data_start_row + len(assignments) + 2
    ws.cell(row=legend_row, column=1, value="Legend:").font = bold
    legend_items = [
    ("A = Morning", MORNING),
    ("B = Evening", EVENING),
    ("C = Night", NIGHT),
    ("WO = Weekly Off", OFF),
    ("L = Leave", LEAVE)
]
    for i, (text, code) in enumerate(legend_items):
        cell = ws.cell(row=legend_row + 1 + i, column=1, value=text)
        cell.fill = PatternFill("solid", fgColor=SHIFT_COLORS[code])

    # Warnings sheet (only added if there are any understaffed days)
    if warnings:
        ws2 = wb.create_sheet("Warnings")
        ws2.cell(row=1, column=1, value="Staffing warnings").font = bold
        for i, w in enumerate(warnings, start=2):
            ws2.cell(row=i, column=1, value=w)
        ws2.column_dimensions["A"].width = 90

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


if __name__ == "__main__":
    # debug=True auto-reloads the site whenever you save a code change,
    # which is handy while you're still building/testing it.
    app.run(debug=True, port=5000)
